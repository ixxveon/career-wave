from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import boto3
from botocore.exceptions import BotoCoreError, ClientError
import pdfplumber
from docx import Document

from admin.ai_metrics.config.settings import get_ai_metrics_settings
from admin.ai_metrics.exception import AiMetricsErrorCode, AiMetricsException
from admin.ai_metrics.repository.rag_document_repository import RagDocumentRecord


@dataclass(frozen=True)
class ParsedRagDocument:
    rag_document_id: int
    original_file_name: str
    mime_type: str | None
    file_path: str
    text: str


class RagDocumentParser:
    """RAG document parsing pipeline entry point."""

    def parse(self, rag_document: RagDocumentRecord) -> ParsedRagDocument:
        document_bytes = self.load_document_bytes(rag_document)
        try:
            text = self._extract_text(
                document_bytes=document_bytes,
                file_path=rag_document.file_path,
                mime_type=rag_document.mime_type,
            )
        except AiMetricsException:
            raise
        except Exception as error:
            raise self._build_indexing_failed_exception(
                rag_document=rag_document,
                reason="document_parse_failed",
            ) from error

        return ParsedRagDocument(
            rag_document_id=rag_document.rag_document_id,
            original_file_name=rag_document.original_file_name,
            mime_type=rag_document.mime_type,
            file_path=rag_document.file_path,
            text=text,
        )

    def load_document_bytes(self, rag_document: RagDocumentRecord) -> bytes:
        settings = get_ai_metrics_settings()

        if settings.file_storage_provider.lower() == "s3":
            return self._load_from_s3(rag_document.file_path)

        resolved_path = self._resolve_local_path(rag_document.file_path)
        try:
            return resolved_path.read_bytes()
        except OSError as error:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={
                    "ragDocumentId": rag_document.rag_document_id,
                    "filePath": rag_document.file_path,
                    "reason": "file_load_failed",
                },
            ) from error

    def _resolve_local_path(self, file_path: str) -> Path:
        settings = get_ai_metrics_settings()
        base_path = Path(settings.file_storage_base_path).resolve()
        candidate = Path(file_path)
        raw_path = candidate if candidate.is_absolute() else (base_path / candidate)
        resolved_path = raw_path.resolve()

        try:
            resolved_path.relative_to(base_path)
        except ValueError as error:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={
                    "filePath": file_path,
                    "reason": "file_path_outside_storage",
                },
            ) from error

        if not resolved_path.is_file():
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={
                    "filePath": file_path,
                    "reason": "file_not_found",
                },
            )

        return resolved_path

    def _load_from_s3(self, file_path: str) -> bytes:
        settings = get_ai_metrics_settings()
        bucket = settings.aws_s3_bucket
        key = file_path.lstrip("/")

        if not bucket:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={
                    "filePath": file_path,
                    "reason": "missing_s3_bucket",
                },
            )

        client = boto3.client(
            "s3",
            aws_access_key_id=settings.aws_access_key_id,
            aws_secret_access_key=settings.aws_secret_access_key,
            region_name=settings.aws_region,
        )

        try:
            response = client.get_object(Bucket=bucket, Key=key)
            return response["Body"].read()
        except (ClientError, BotoCoreError) as error:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={
                    "filePath": file_path,
                    "reason": "file_load_failed",
                },
            ) from error

    def _extract_text(
        self,
        *,
        document_bytes: bytes,
        file_path: str,
        mime_type: str | None,
    ) -> str:
        suffix = Path(file_path).suffix.lower()

        if mime_type == "application/pdf" or suffix == ".pdf":
            return self._extract_pdf_text(document_bytes)

        if (
            mime_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or suffix == ".docx"
        ):
            return self._extract_docx_text(document_bytes)

        if mime_type is not None and mime_type.startswith("text/"):
            return self._extract_plain_text(document_bytes)

        raise AiMetricsException(
            error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
            detail={
                "filePath": file_path,
                "mimeType": mime_type,
                "reason": "unsupported_document_format",
            },
        )

    def _extract_pdf_text(self, document_bytes: bytes) -> str:
        with pdfplumber.open(BytesIO(document_bytes)) as pdf:
            pages_text = [page.extract_text() or "" for page in pdf.pages]

        text = self._normalize_text("\n".join(pages_text))
        if not text:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={"reason": "empty_pdf_text"},
            )
        return text

    def _extract_docx_text(self, document_bytes: bytes) -> str:
        document = Document(BytesIO(document_bytes))
        text = self._normalize_text(
            "\n".join(paragraph.text for paragraph in document.paragraphs),
        )
        if not text:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={"reason": "empty_docx_text"},
            )
        return text

    def _extract_plain_text(self, document_bytes: bytes) -> str:
        try:
            text = self._normalize_text(document_bytes.decode("utf-8"))
        except UnicodeDecodeError as error:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={"reason": "plain_text_decode_failed"},
            ) from error

        if not text:
            raise AiMetricsException(
                error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
                detail={"reason": "empty_plain_text"},
            )

        return text

    def _normalize_text(self, text: str) -> str:
        return text.replace("\r\n", "\n").replace("\r", "\n").strip()

    def _build_indexing_failed_exception(
        self,
        *,
        rag_document: RagDocumentRecord,
        reason: str,
    ) -> AiMetricsException:
        return AiMetricsException(
            error_code=AiMetricsErrorCode.RAG_DOCUMENT_INDEXING_FAILED,
            detail={
                "ragDocumentId": rag_document.rag_document_id,
                "filePath": rag_document.file_path,
                "mimeType": rag_document.mime_type,
                "reason": reason,
            },
        )
